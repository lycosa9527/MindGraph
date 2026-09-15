"""Resolve, inspect, and manage MindMate teaching-design Word templates."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from docx import Document
from docx.document import Document as DocumentType
from docx.opc.exceptions import PackageNotFoundError

from services.mindmate.teaching_design_template_catalog import (
    add_catalog_template,
    default_template_id,
    delete_catalog_template,
    find_template,
    list_known_template_ids,
    load_catalog,
    set_default_bundled,
    template_exists_on_disk,
    template_file_path,
    update_catalog_template,
)
from services.mindmate.teaching_design_template_preview import invalidate_template_preview
from services.mindmate.teaching_design_template_paths import (
    BUNDLED_TEMPLATE_PATH,
    DOCX_MAGIC,
    MAX_TEMPLATE_BYTES,
    TEMPLATE_KEY_BUNDLED,
    TEMPLATE_KEY_SYSTEM,
    TEMPLATE_KEY_UPLOADED,
)
from services.utils.error_types import FILE_IO_ERRORS

DOCX_READ_ERRORS = (*FILE_IO_ERRORS, BadZipFile, PackageNotFoundError, KeyError)


def resolve_active_template_path() -> Path:
    """Word file used when a school follows the system default."""
    return template_file_path(default_template_id())


def resolve_template_path(template_key: str | None) -> Path:
    """Resolve a school-selected key, or follow the current system template."""
    key = (template_key or "").strip()
    if key in {"", TEMPLATE_KEY_SYSTEM}:
        return resolve_active_template_path()
    return template_file_path(key)


def normalize_template_key(value: object) -> str | None:
    """Store None for 跟随系统; reject unknown or missing uploaded pins."""
    if value is None:
        return None
    key = str(value).strip()
    if key in {"", TEMPLATE_KEY_SYSTEM}:
        return None
    if key not in list_known_template_ids():
        raise ValueError("teaching_design_template_invalid_key")
    if not template_exists_on_disk(key):
        raise ValueError("teaching_design_template_uploaded_missing")
    return key


def teaching_design_template_list_field(org: object) -> dict[str, str | None]:
    """API field for organization list/detail payloads."""
    raw = getattr(org, "teaching_design_template_key", None)
    if isinstance(raw, str) and raw.strip():
        return {"teaching_design_template_key": raw.strip()}
    return {"teaching_design_template_key": None}


def list_template_options() -> list[dict[str, str]]:
    """Keys the school modal can pin, including named catalog rows."""
    state = load_catalog()
    options = [
        {"key": TEMPLATE_KEY_SYSTEM, "source": TEMPLATE_KEY_SYSTEM, "filename": "", "name": ""},
        {
            "key": TEMPLATE_KEY_BUNDLED,
            "source": TEMPLATE_KEY_BUNDLED,
            "filename": BUNDLED_TEMPLATE_PATH.name,
            "name": str(state.get("bundled_name") or ""),
        },
    ]
    for item in state["templates"]:
        path = template_file_path(str(item["id"]))
        if not path.is_file() or path == BUNDLED_TEMPLATE_PATH:
            continue
        options.append(
            {
                "key": str(item["id"]),
                "source": TEMPLATE_KEY_UPLOADED,
                "filename": str(item.get("original_filename") or ""),
                "name": str(item.get("name") or ""),
            }
        )
    return options


def inspect_document_tables(document: DocumentType) -> list[dict[str, Any]]:
    """Summarize Word tables for the admin edit preview."""
    rows: list[dict[str, Any]] = []
    for index, table in enumerate(document.tables, start=1):
        first_cells: list[str] = []
        if table.rows:
            first_cells = [cell.text.strip() for cell in table.rows[0].cells]
        preview = " / ".join(part for part in first_cells if part)
        rows.append(
            {
                "index": index,
                "row_count": len(table.rows),
                "col_count": len(first_cells),
                "preview": preview[:160],
            }
        )
    return rows


def inspect_docx_tables(path: Path) -> list[dict[str, Any]]:
    """Open a .docx path and return table summaries."""
    return inspect_document_tables(Document(str(path)))


def validate_docx_bytes(payload: bytes) -> None:
    """Reject non-DOCX uploads before they become catalog files."""
    if len(payload) > MAX_TEMPLATE_BYTES:
        raise ValueError("teaching_design_template_too_large")
    if len(payload) < 4 or not payload.startswith(DOCX_MAGIC):
        raise ValueError("teaching_design_template_invalid")
    try:
        inspect_document_tables(Document(io.BytesIO(payload)))
    except DOCX_READ_ERRORS as exc:
        raise ValueError("teaching_design_template_invalid") from exc


def _mtime_iso(path: Path) -> str:
    """ISO timestamp from the file mtime."""
    return datetime.fromtimestamp(path.stat().st_mtime, tz=UTC).isoformat()


def _bundled_row(state: dict[str, Any], default_id: str) -> dict[str, Any]:
    """Landing row for the committed BNU form."""
    path = BUNDLED_TEMPLATE_PATH
    return {
        "id": TEMPLATE_KEY_BUNDLED,
        "index": 1,
        "name": str(state.get("bundled_name") or ""),
        "source": TEMPLATE_KEY_BUNDLED,
        "filename": path.name,
        "size_bytes": path.stat().st_size if path.is_file() else 0,
        "updated_at": _mtime_iso(path) if path.is_file() else None,
        "uploaded_by_name": None,
        "is_default": default_id == TEMPLATE_KEY_BUNDLED,
        "can_delete": False,
    }


def _uploaded_row(item: dict[str, Any], index: int, default_id: str) -> dict[str, Any]:
    """Landing row for one uploaded catalog template."""
    path = template_file_path(str(item["id"]))
    updated = item.get("uploaded_at")
    if not isinstance(updated, str) or not updated.strip():
        updated = _mtime_iso(path) if path.is_file() else None
    return {
        "id": item["id"],
        "index": index,
        "name": str(item.get("name") or ""),
        "source": TEMPLATE_KEY_UPLOADED,
        "filename": str(item.get("original_filename") or ""),
        "size_bytes": int(item.get("size_bytes") or 0),
        "updated_at": updated,
        "uploaded_by_name": item.get("uploaded_by_name"),
        "is_default": default_id == item["id"],
        "can_delete": True,
    }


def describe_template_catalog() -> dict[str, Any]:
    """Landing-page rows: bundled first, then named uploads."""
    state = load_catalog()
    default_id = str(state.get("default_id") or TEMPLATE_KEY_BUNDLED)
    rows = [_bundled_row(state, default_id)]
    for offset, item in enumerate(state["templates"]):
        rows.append(_uploaded_row(item, offset + 2, default_id))
    return {
        "default_id": default_id,
        "can_restore": default_id != TEMPLATE_KEY_BUNDLED,
        "templates": rows,
    }


def describe_active_template() -> dict[str, Any]:
    """Metadata plus Word tables for the current system-default file."""
    path = resolve_active_template_path()
    if not path.is_file():
        raise FileNotFoundError(f"Teaching-design template missing: {path}")
    catalog = describe_template_catalog()
    default_id = str(catalog.get("default_id") or TEMPLATE_KEY_BUNDLED)
    current = next(
        (row for row in catalog["templates"] if row["id"] == default_id),
        catalog["templates"][0],
    )
    return {
        "filename": current["filename"],
        "source": current["source"],
        "size_bytes": current["size_bytes"],
        "updated_at": current["updated_at"],
        "uploaded_by_name": current["uploaded_by_name"],
        "can_restore": catalog["can_restore"],
        "tables": inspect_docx_tables(path),
    }


def save_uploaded_template(
    payload: bytes,
    *,
    original_filename: str,
    uploaded_by_user_id: int | None,
    uploaded_by_name: str | None,
    display_name: str | None = None,
) -> dict[str, Any]:
    """Add a named template to the catalog and return the landing payload."""
    validate_docx_bytes(payload)
    add_catalog_template(
        payload,
        original_filename=original_filename,
        display_name=display_name,
        uploaded_by_user_id=uploaded_by_user_id,
        uploaded_by_name=uploaded_by_name,
    )
    return describe_template_catalog()


def update_named_template(
    template_id: str,
    *,
    name: str | None = None,
    is_default: bool | None = None,
    payload: bytes | None = None,
    original_filename: str | None = None,
    uploaded_by_user_id: int | None = None,
    uploaded_by_name: str | None = None,
) -> dict[str, Any]:
    """Edit name / default / file, then return the landing payload."""
    if payload is not None:
        validate_docx_bytes(payload)
    update_catalog_template(
        template_id,
        name=name,
        is_default=is_default,
        payload=payload,
        original_filename=original_filename,
        uploaded_by_user_id=uploaded_by_user_id,
        uploaded_by_name=uploaded_by_name,
    )
    if payload is not None:
        invalidate_template_preview(template_id)
    return describe_template_catalog()


def delete_named_template(template_id: str) -> dict[str, Any]:
    """Delete an uploaded catalog row and return the landing payload."""
    delete_catalog_template(template_id)
    invalidate_template_preview(template_id)
    return describe_template_catalog()


def restore_bundled_template(*, user_id: int | None) -> dict[str, Any]:
    """Point the system default at the bundled form; keep uploaded rows."""
    set_default_bundled(user_id=user_id)
    return describe_template_catalog()


def resolve_download_path(template_id: str) -> tuple[Path, str]:
    """Path and download filename for a catalog row."""
    key = template_id.strip()
    if key == TEMPLATE_KEY_BUNDLED:
        return BUNDLED_TEMPLATE_PATH, BUNDLED_TEMPLATE_PATH.name
    item = find_template(key)
    if item is None:
        raise FileNotFoundError(f"Teaching-design template missing: {key}")
    path = template_file_path(key)
    if not path.is_file():
        raise FileNotFoundError(f"Teaching-design template missing: {key}")
    filename = str(item.get("original_filename") or path.name)
    return path, Path(filename).name

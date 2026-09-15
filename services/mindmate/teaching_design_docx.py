"""Fill the committed BNU thinking-classroom teaching-design DOCX template."""

from __future__ import annotations

import io

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph

from services.mindmate.teaching_design_models import (
    ActivityStage,
    TeachingDesignSpec,
    ThinkingPointRow,
)
from services.mindmate.teaching_design_template_store import (
    BUNDLED_TEMPLATE_PATH,
    resolve_template_path,
)

TEMPLATE_PATH = BUNDLED_TEMPLATE_PATH

_TITLE_PLACEHOLDER = "“”/《   》教学设计"
_HEADING_PREFIXES = (
    "【摘要】",
    "【思维训练点】",
    "【教学内容分析】",
    "【学习者分析】",
    "【学习目标及重难点】",
    "【课例结构】",
    "【学习活动设计】",
    "【板书设计】",
    "【作业与拓展学习设计】",
    "【素材设计】",
)
_THINKING_SLOTS = {
    "问题情境": (1, 4),
    "认知冲突": (4, 7),
    "思维可视化": (7, 9),
    "变式运用": (9, 11),
}
_ACTIVITY_BLOCKS = ((2, 5), (6, 9), (10, 13))


def build_teaching_design_docx(
    spec: TeachingDesignSpec,
    *,
    template_key: str | None = None,
) -> bytes:
    """Return filled DOCX bytes from the school-selected or system template."""
    template_path = resolve_template_path(template_key)
    if not template_path.is_file():
        raise FileNotFoundError(f"Teaching-design template missing: {template_path}")
    document = Document(str(template_path))
    _fill_title(document, spec)
    if document.tables:
        _fill_header_table(document.tables[0], spec)
    if len(document.tables) > 1:
        _fill_thinking_table(document.tables[1], spec.thinking_points)
    extra_stages: list[ActivityStage] = []
    if len(document.tables) > 2:
        extra_stages = _fill_activity_table(document.tables[2], spec.activities)
    _replace_section_body(document, "【摘要】", spec.summary)
    _replace_section_body(document, "【教学内容分析】", spec.content_analysis)
    _replace_section_body(document, "【学习者分析】", spec.learner_analysis)
    _replace_section_body(document, "【学习目标及重难点】", spec.objectives)
    _replace_section_body(document, "【课例结构】", spec.structure)
    overflow = _format_extra_stages(extra_stages)
    _replace_section_body(document, "【学习活动设计】", overflow)
    _replace_section_body(document, "【板书设计】", spec.board_design)
    _replace_section_body(document, "【作业与拓展学习设计】", spec.homework)
    _replace_section_body(document, "【素材设计】", spec.materials)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fill_title(document: DocumentType, spec: TeachingDesignSpec) -> None:
    title = spec.title.strip() or _TITLE_PLACEHOLDER
    if not title.endswith("教学设计") and title:
        title = f"{title}教学设计"
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == _TITLE_PLACEHOLDER:
            _write_paragraph(paragraph, title)
            return


def _fill_header_table(table: Table, spec: TeachingDesignSpec) -> None:
    _write_cell(table.cell(0, 1), spec.grade or table.cell(0, 1).text)
    _write_cell(table.cell(0, 3), spec.subject or table.cell(0, 3).text)
    _write_cell(table.cell(1, 1), spec.textbook or table.cell(1, 1).text)
    _write_cell(table.cell(1, 3), spec.period or table.cell(1, 3).text)
    school = spec.school.strip() or table.cell(2, 1).text
    teacher = spec.teacher.strip() or table.cell(2, 3).text
    _write_cell(table.cell(2, 1), school)
    _write_cell(table.cell(2, 3), teacher)


def _fill_thinking_table(table: Table, rows: list[ThinkingPointRow]) -> None:
    grouped: dict[str, list[ThinkingPointRow]] = {key: [] for key in _THINKING_SLOTS}
    leftovers: list[ThinkingPointRow] = []
    for row in rows:
        category = row.category.strip()
        if category in grouped:
            grouped[category].append(row)
        else:
            leftovers.append(row)
    if leftovers:
        grouped["问题情境"].extend(leftovers)
    for category, (start, end) in _THINKING_SLOTS.items():
        slots = grouped[category]
        for offset, row_index in enumerate(range(start, end)):
            if offset >= len(slots):
                break
            item = slots[offset]
            point = item.point.strip() or table.cell(row_index, 1).text
            _write_cell(table.cell(row_index, 1), point)
            if item.intent.strip():
                _write_cell(table.cell(row_index, 2), item.intent.strip())


def _fill_activity_table(table: Table, stages: list[ActivityStage]) -> list[ActivityStage]:
    for index, (title_row, intent_row) in enumerate(_ACTIVITY_BLOCKS):
        if index >= len(stages):
            break
        stage = stages[index]
        title = stage.title.strip()
        _write_cell(table.cell(title_row, 0), f"环节{index + 1}：{title}".rstrip("："))
        teacher_one = stage.teacher_actions[0] if stage.teacher_actions else ""
        teacher_two = "\n".join(stage.teacher_actions[1:]) if len(stage.teacher_actions) > 1 else ""
        student_one = stage.student_actions[0] if stage.student_actions else ""
        student_two = "\n".join(stage.student_actions[1:]) if len(stage.student_actions) > 1 else ""
        if teacher_one:
            _write_cell(table.cell(title_row + 1, 0), f"教师活动1：{teacher_one}")
        if student_one:
            _write_cell(table.cell(title_row + 1, 1), f"学生活动1：{student_one}")
        if teacher_two:
            _write_cell(table.cell(title_row + 2, 0), f"教师活动2：{teacher_two}")
        if student_two:
            _write_cell(table.cell(title_row + 2, 1), f"学生活动2：{student_two}")
        if stage.intent.strip():
            _write_cell(table.cell(intent_row, 0), f"设计意图说明：{stage.intent.strip()}")
    return stages[len(_ACTIVITY_BLOCKS) :]


def _format_extra_stages(stages: list[ActivityStage]) -> str:
    blocks: list[str] = []
    for offset, stage in enumerate(stages, start=4):
        lines = [f"环节{offset}：{stage.title}".rstrip("：")]
        for action in stage.teacher_actions:
            lines.append(f"教师活动：{action}")
        for action in stage.student_actions:
            lines.append(f"学生活动：{action}")
        if stage.intent.strip():
            lines.append(f"设计意图：{stage.intent.strip()}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def _replace_section_body(document: DocumentType, heading: str, text: str) -> None:
    content = (text or "").strip()
    if not content:
        return
    paragraphs = document.paragraphs
    start = _find_heading_index(paragraphs, heading)
    if start is None:
        return
    body_indexes = _section_body_indexes(paragraphs, start)
    if not body_indexes:
        return
    lines = content.splitlines() or [content]
    _write_paragraph(paragraphs[body_indexes[0]], lines[0])
    extra_index = 1
    for paragraph_index in body_indexes[1:]:
        if extra_index < len(lines):
            _write_paragraph(paragraphs[paragraph_index], lines[extra_index])
            extra_index += 1
            continue
        if paragraphs[paragraph_index].text.strip():
            _write_paragraph(paragraphs[paragraph_index], "")
    if extra_index < len(lines):
        tail = "\n".join(lines[extra_index:])
        current = paragraphs[body_indexes[0]].text
        _write_paragraph(paragraphs[body_indexes[0]], f"{current}\n{tail}" if current else tail)


def _find_heading_index(paragraphs: list[Paragraph], heading: str) -> int | None:
    prefix = heading.rstrip()
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip().startswith(prefix):
            return index
    return None


def _section_body_indexes(paragraphs: list[Paragraph], heading_index: int) -> list[int]:
    indexes: list[int] = []
    for index in range(heading_index + 1, len(paragraphs)):
        text = paragraphs[index].text.strip()
        if text.startswith("附录"):
            break
        if any(text.startswith(prefix) for prefix in _HEADING_PREFIXES):
            break
        if text.startswith("说明："):
            continue
        indexes.append(index)
    return indexes


def _write_cell(cell, text: str) -> None:
    cell.text = text


def _write_paragraph(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    paragraph.text = text

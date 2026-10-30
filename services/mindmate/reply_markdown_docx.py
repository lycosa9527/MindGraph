"""Turn a MindMate reply into a Word document with markdown and math rendered."""

from __future__ import annotations

import io
import re

from docx import Document as create_document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from services.mindmate.reply_markdown_inline import (
    InlineSpan,
    clean_docx_text,
    inline_plain_text,
    linearize_latex,
    parse_inline,
)

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_LIST_RE = re.compile(r"^([ \t]*)([-+*]|\d{1,3}[.)])\s+(.*?)\s*$")
_HR_RE = re.compile(r"^([-*_])\1{2,}\s*$")
_TABLE_ROW_RE = re.compile(r"^\s*\|.+\|\s*$")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
_BOLD_TITLE_RE = re.compile(r"^\*\*(.+)\*\*$")
_SAFE_NAME = re.compile(r"[^\w\-.\u4e00-\u9fff]+", re.UNICODE)
_MATH_FONT = "Cambria Math"
_CODE_FONT = "Consolas"
_ORDERED_HANG_CM = 1.0
_BULLET_HANG_CM = 0.6


def sanitize_docx_stem(name: str, fallback: str) -> str:
    """Keep a short filename stem that is safe on disk and in headers."""
    cleaned = _SAFE_NAME.sub("-", (name or "").strip())[:80].strip("-.")
    return cleaned or fallback


def reply_docx_filename(markdown: str, fallback: str) -> str:
    """Name the file from the first heading, then the first line of text."""
    for line in (markdown or "").splitlines():
        heading = _HEADING_RE.match(line.strip())
        if heading:
            return f"{sanitize_docx_stem(inline_plain_text(heading.group(2)), fallback)}.docx"
    for line in (markdown or "").splitlines():
        plain = inline_plain_text(line.strip())
        if plain.strip():
            return f"{sanitize_docx_stem(plain, fallback)}.docx"
    return f"{fallback}.docx"


def build_markdown_docx(markdown: str) -> bytes:
    """Write headings, lists, tables, and formulas into a new docx."""
    document = create_document()
    lines = (markdown or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    index = 0
    while index < len(lines):
        if not lines[index].strip():
            index += 1
            continue
        next_index = _write_block(document, lines, index)
        if next_index <= index:
            index += 1
            continue
        index = next_index
    if not _document_has_text(document):
        raise ValueError("empty_reply")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _write_block(document: DocumentType, lines: list[str], index: int) -> int:
    line = lines[index]
    if _opening_fence(line):
        return _write_fence(document, lines, index)
    if _is_table(lines, index):
        return _write_table(document, lines, index)
    heading = _HEADING_RE.match(line.strip())
    if heading:
        _write_heading(document, len(heading.group(1)), heading.group(2).strip())
        return index + 1
    if _HR_RE.match(line.strip()):
        _write_rule(document)
        return index + 1
    if line.lstrip().startswith(">"):
        return _write_quote(document, lines, index)
    if _LIST_RE.match(line):
        return _write_list(document, lines, index)
    title = _bold_title(line)
    if title is not None:
        _write_heading(document, 2, title)
        return index + 1
    display_end = _write_display_math(document, lines, index)
    if display_end is not None:
        return display_end
    _write_paragraph(document, line.strip())
    return index + 1


def _write_heading(document: DocumentType, level: int, text: str) -> None:
    style = f"Heading {min(max(level, 1), 6)}"
    paragraph = document.add_paragraph(style=style)
    _write_spans(paragraph, parse_inline(text))


def _write_paragraph(document: DocumentType, text: str, *, quote: bool = False) -> None:
    paragraph = document.add_paragraph()
    if quote:
        paragraph.paragraph_format.left_indent = Cm(0.75)
    spans = parse_inline(text, italic=quote)
    if _is_display_math(text, spans):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _write_spans(paragraph, spans)


def _write_quote(document: DocumentType, lines: list[str], index: int) -> int:
    while index < len(lines) and lines[index].lstrip().startswith(">"):
        body = re.sub(r"^\s*>\s?", "", lines[index]).strip()
        if body:
            _write_paragraph(document, body, quote=True)
        index += 1
    return index


def _write_list(document: DocumentType, lines: list[str], index: int) -> int:
    while index < len(lines):
        match = _LIST_RE.match(lines[index])
        if match is None:
            break
        indent, marker, body = match.group(1), match.group(2), match.group(3)
        index += 1
        extras: list[str] = []
        while index < len(lines) and _is_list_continuation(lines[index]):
            extras.append(lines[index].strip())
            index += 1
        _write_list_item(document, indent, marker, body, extras)
    return index


def _write_list_item(
    document: DocumentType,
    indent: str,
    marker: str,
    body: str,
    extras: list[str],
) -> None:
    ordered = marker[:1].isdigit()
    task, rest = _split_task(body)
    level = min(len(indent.replace("\t", "  ")) // 2, 6)
    hang = _ORDERED_HANG_CM if ordered else _BULLET_HANG_CM
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(hang + 0.75 * level)
    paragraph.paragraph_format.first_line_indent = Cm(-hang)
    if task and ordered:
        prefix = f"{marker} {task.strip()} "
    elif task:
        prefix = task
    elif ordered:
        prefix = f"{marker} "
    else:
        prefix = "• "
    paragraph.add_run(clean_docx_text(prefix))
    _write_spans(paragraph, parse_inline(rest))
    for extra in extras:
        paragraph.add_run().add_break()
        _write_spans(paragraph, parse_inline(extra))


def _write_fence(document: DocumentType, lines: list[str], index: int) -> int:
    marker = _opening_fence(lines[index])
    if marker is None:
        return index + 1
    index += 1
    while index < len(lines) and not _closing_fence(lines[index], marker):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(clean_docx_text(lines[index]))
        run.font.name = _CODE_FONT
        run.font.size = Pt(10)
        index += 1
    if index < len(lines):
        index += 1
    return index


def _write_table(document: DocumentType, lines: list[str], index: int) -> int:
    rows = [_split_cells(lines[index])]
    index += 2
    while index < len(lines) and _TABLE_ROW_RE.match(lines[index]):
        rows.append(_split_cells(lines[index]))
        index += 1
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            source = row[col_index] if col_index < len(row) else ""
            paragraph = table.cell(row_index, col_index).paragraphs[0]
            _write_spans(paragraph, parse_inline(source, bold=row_index == 0))
    return index


def _write_rule(document: DocumentType) -> None:
    document.add_paragraph()
    paragraph = document.element.body.p_lst[-1]
    properties = paragraph.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    borders.append(bottom)
    properties.append(borders)


def _write_spans(paragraph: Paragraph, spans: list[InlineSpan]) -> None:
    if not spans:
        if not paragraph.runs:
            paragraph.add_run("")
        return
    for span in spans:
        run = paragraph.add_run(clean_docx_text(span.text))
        if span.bold:
            run.bold = True
        if span.italic or span.math:
            run.italic = True
        if span.strike:
            run.font.strike = True
        if span.code:
            run.font.name = _CODE_FONT
        elif span.math:
            run.font.name = _MATH_FONT
        if span.link and span.link not in span.text:
            paragraph.add_run(clean_docx_text(f" ({span.link})"))


def _is_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    return _TABLE_ROW_RE.match(lines[index]) is not None and _TABLE_SEP_RE.match(lines[index + 1]) is not None


def _split_cells(line: str) -> list[str]:
    """Split a markdown row, keeping `$|a|$` inside one cell."""
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    cells: list[str] = []
    buffer: list[str] = []
    index = 0
    while index < len(body):
        if body.startswith("$$", index):
            end = body.find("$$", index + 2)
            if end == -1:
                buffer.append(body[index:])
                break
            buffer.append(body[index : end + 2])
            index = end + 2
            continue
        if body[index] == "$":
            end = body.find("$", index + 1)
            if end != -1:
                buffer.append(body[index : end + 1])
                index = end + 1
                continue
        if body[index] == "|" and (index == 0 or body[index - 1] != "\\"):
            cells.append("".join(buffer).strip())
            buffer = []
            index += 1
            continue
        buffer.append(body[index])
        index += 1
    cells.append("".join(buffer).strip())
    return cells


def _is_list_continuation(line: str) -> bool:
    if not line.strip() or _LIST_RE.match(line):
        return False
    if _HEADING_RE.match(line.strip()) or _opening_fence(line) or line.lstrip().startswith(">"):
        return False
    return line.startswith((" ", "\t"))


def _bold_title(line: str) -> str | None:
    match = _BOLD_TITLE_RE.fullmatch(line.strip())
    if match is None:
        return None
    inner = match.group(1).strip()
    if not inner or "**" in inner or len(inner) > 80:
        return None
    return inner


def _is_display_math(text: str, spans: list[InlineSpan]) -> bool:
    stripped = text.strip()
    display = stripped.startswith("$$") or stripped.startswith(r"\[")
    return display and len(spans) == 1 and spans[0].math


def _write_display_math(document: DocumentType, lines: list[str], index: int) -> int | None:
    """Render a multi-line `$$` or `\\[` block. A one-line formula stays inline."""
    stripped = lines[index].strip()
    pair = _display_delimiters(stripped)
    if pair is None or _closed_on_same_line(stripped, pair):
        return None
    opener, closer = pair
    body: list[str] = []
    if stripped != opener:
        body.append(stripped[len(opener) :].strip())
    cursor = index + 1
    while cursor < len(lines):
        current = lines[cursor].strip()
        if current == closer or current.endswith(closer):
            piece = current[: -len(closer)].strip()
            if piece:
                body.append(piece)
            _write_math_paragraph(document, "\n".join(body))
            return cursor + 1
        body.append(lines[cursor])
        cursor += 1
    return None


def _display_delimiters(stripped: str) -> tuple[str, str] | None:
    if stripped.startswith("$$"):
        return "$$", "$$"
    if stripped.startswith(r"\["):
        return r"\[", r"\]"
    return None


def _closed_on_same_line(stripped: str, pair: tuple[str, str]) -> bool:
    opener, closer = pair
    if stripped == opener or len(stripped) < len(opener) + len(closer):
        return False
    return stripped.startswith(opener) and stripped.endswith(closer)


def _write_math_paragraph(document: DocumentType, body: str) -> None:
    linear = linearize_latex(body)
    if not linear:
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _write_spans(paragraph, [InlineSpan(linear, math=True)])


def _document_has_text(document: DocumentType) -> bool:
    if any(paragraph.text.strip() for paragraph in document.paragraphs):
        return True
    return bool(document.tables)


def _opening_fence(line: str) -> str | None:
    stripped = line.lstrip()
    if stripped.startswith("```"):
        char = "`"
    elif stripped.startswith("~~~"):
        char = "~"
    else:
        return None
    count = 0
    for item in stripped:
        if item != char:
            break
        count += 1
    if count < 3:
        return None
    return char * count


def _closing_fence(line: str, opener: str) -> bool:
    stripped = line.strip()
    return bool(stripped) and len(stripped) >= len(opener) and set(stripped) == {opener[0]}


def _split_task(body: str) -> tuple[str, str]:
    if body.startswith("[ ] "):
        return "[ ] ", body[4:]
    if body.startswith("[x] ") or body.startswith("[X] "):
        return "[x] ", body[4:]
    return "", body

"""OOXML / CSV text extraction helpers for DocumentProcessor.

Author: lycosa9527
Made by: MindSpring Team

Copyright 2024-2025 北京思源智教科技有限公司 (Beijing Siyuan Zhijiao Technology Co., Ltd.)
All Rights Reserved
Proprietary License
"""

from __future__ import annotations

import csv
from typing import Any, List, Optional

from services.knowledge.markdown_tables import rows_to_markdown_table, sheet_to_markdown
from services.utils.error_types import FILE_IO_ERRORS


class _DocxOxmlApi:
    """Optional python-docx OXMl helpers for ordered paragraph/table walks."""

    qn: Any = None
    table: Any = None
    paragraph: Any = None
    available: bool = False


_DOCX_OXML = _DocxOxmlApi()
try:
    from docx.oxml.ns import qn as _docx_qn_import
    from docx.table import Table as _docx_table_import
    from docx.text.paragraph import Paragraph as _docx_paragraph_import

    _DOCX_OXML.qn = _docx_qn_import
    _DOCX_OXML.table = _docx_table_import
    _DOCX_OXML.paragraph = _docx_paragraph_import
    _DOCX_OXML.available = True
except ImportError:
    pass


def extract_docx_markdown(document_cls: Any, file_path: str) -> str:
    """Extract paragraphs and tables from a DOCX as markdown."""
    doc = document_cls(file_path)
    parts: List[str] = []

    body = getattr(getattr(doc, "element", None), "body", None)
    if body is not None and _DOCX_OXML.available:
        for child in body.iterchildren():
            tag = child.tag
            if tag == _DOCX_OXML.qn("w:p"):
                paragraph = _DOCX_OXML.paragraph(child, doc)
                text = paragraph.text.strip()
                if text:
                    parts.append(text)
            elif tag == _DOCX_OXML.qn("w:tbl"):
                table = _DOCX_OXML.table(child, doc)
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                md_table = rows_to_markdown_table(rows)
                if md_table:
                    parts.append(md_table)
        if parts:
            return "\n\n".join(parts)

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        md_table = rows_to_markdown_table(rows)
        if md_table:
            parts.append(md_table)
    return "\n\n".join(parts)


def extract_pptx_markdown(presentation_cls: Any, file_path: str) -> str:
    """Extract slide shape text and speaker notes as markdown."""
    presentation = presentation_cls(file_path)
    parts: List[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts: List[str] = []
        for shape in slide.shapes:
            shape_text = getattr(shape, "text", None)
            if shape_text and str(shape_text).strip():
                slide_parts.append(str(shape_text).strip())
        notes_text = _slide_notes_text(slide)
        if notes_text:
            slide_parts.append(f"*Notes:* {notes_text}")
        if slide_parts:
            parts.append(f"## Slide {index}\n\n" + "\n\n".join(slide_parts))
    return "\n\n".join(parts)


def _slide_notes_text(slide: Any) -> str:
    """Return speaker notes for a slide, if present."""
    try:
        notes_slide = slide.notes_slide
    except (AttributeError, ValueError, KeyError):
        return ""
    if notes_slide is None:
        return ""
    frame = getattr(notes_slide, "notes_text_frame", None)
    if frame is None:
        return ""
    text = getattr(frame, "text", "") or ""
    return str(text).strip()


def extract_xlsx_markdown(load_workbook: Any, file_path: str) -> str:
    """Extract workbook sheets as markdown tables."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    parts: List[str] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            section = sheet_to_markdown(str(sheet_name), rows)
            if section:
                parts.append(section)
    finally:
        workbook.close()
    return "\n\n".join(parts)


def extract_csv_markdown(file_path: str) -> str:
    """Parse a CSV file into a markdown table."""
    encodings = ("utf-8-sig", "utf-8", "gbk", "gb2312", "latin-1")
    last_error: Optional[Exception] = None
    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding, newline="") as handle:
                sample = handle.read(4096)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.reader(handle, dialect)
                rows = list(reader)
            table = rows_to_markdown_table(rows)
            if not table:
                raise ValueError("CSV file contains no data rows")
            return table
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
        except FILE_IO_ERRORS as exc:
            last_error = exc
            continue
    raise ValueError(f"Failed to parse CSV: {last_error}")

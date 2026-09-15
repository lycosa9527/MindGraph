"""Fill the committed BNU teaching-design template from a spec."""

from __future__ import annotations

import io

from docx import Document

from routers.api.teaching_design_docx_export import _content_disposition
from services.mindmate.teaching_design_docx import build_teaching_design_docx
from services.mindmate.teaching_design_models import (
    ActivityStage,
    TeachingDesignSpec,
    ThinkingPointRow,
)


def test_build_docx_fills_title_header_and_sections() -> None:
    """Title, header table, 摘要, thinking row, and one activity block are written."""
    spec = TeachingDesignSpec(
        title="《呼吸作用》教学设计",
        grade="初中七年级",
        subject="生物",
        textbook="人教版",
        period="第1课时",
        teacher="王老师",
        summary="用萌发种子的实验冲突重建呼吸作用概念。",
        content_analysis="围绕细胞呼吸展开。",
        learner_analysis="学生已学过光合作用。",
        objectives="1. 说出呼吸作用的概念（重点）",
        thinking_points=[
            ThinkingPointRow(
                category="认知冲突",
                point="萌发种子放热与植物只进行光合作用的矛盾",
                intent="暴露迷思概念",
            )
        ],
        activities=[
            ActivityStage(
                title="创设情境",
                teacher_actions=["出示种子萌发视频"],
                student_actions=["观察并提问"],
                intent="用生活现象引出课题",
            )
        ],
    )
    document = Document(io.BytesIO(build_teaching_design_docx(spec)))
    paragraphs = "\n".join(item.text for item in document.paragraphs)
    assert "《呼吸作用》教学设计" in paragraphs
    assert "用萌发种子的实验冲突" in paragraphs
    assert "围绕细胞呼吸展开" in paragraphs
    header = document.tables[0]
    assert header.cell(0, 1).text == "初中七年级"
    assert header.cell(0, 3).text == "生物"
    assert header.cell(2, 3).text == "王老师"
    thinking = document.tables[1]
    assert "萌发种子放热" in thinking.cell(4, 1).text
    activity = document.tables[2]
    assert "创设情境" in activity.cell(2, 0).text
    assert "出示种子萌发视频" in activity.cell(3, 0).text


def test_content_disposition_allows_chinese_filename() -> None:
    """Content-Disposition stays latin-1 safe and keeps UTF-8 filename*."""
    header = _content_disposition("呼吸作用.docx")
    header.encode("latin-1")
    assert "filename*=UTF-8''" in header

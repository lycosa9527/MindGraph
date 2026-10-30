"""Plain MindMate replies become readable Word files."""

from __future__ import annotations

import io

from docx import Document as open_document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH

from services.mindmate.reply_markdown_docx import (
    build_markdown_docx,
    reply_docx_filename,
)
from services.mindmate.reply_markdown_inline import linearize_latex


def _document_text(markdown: str) -> str:
    payload = build_markdown_docx(markdown)
    document = open_document(io.BytesIO(payload))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    chunks.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(chunks)


def _open(markdown: str) -> DocumentType:
    return open_document(io.BytesIO(build_markdown_docx(markdown)))


def test_markdown_docx_keeps_heading_and_paragraph() -> None:
    """Headings and body text survive the round trip."""
    text = _document_text("# 研讨纪要\n\n先看**问题**，再看证据。\n- 一条")
    assert "研讨纪要" in text
    assert "问题" in text
    assert "一条" in text
    assert "**" not in text


def test_reply_docx_filename_uses_heading() -> None:
    """The artifact name comes from the first heading."""
    assert reply_docx_filename("# 分数的意义\n正文", "研讨文档") == "分数的意义.docx"
    assert reply_docx_filename("   ", "研讨文档") == "研讨文档.docx"


def test_empty_markdown_is_rejected() -> None:
    """A blank reply is not archived as an empty document."""
    try:
        build_markdown_docx("\n\n")
    except ValueError as exc:
        assert str(exc) == "empty_reply"
    else:
        raise AssertionError("expected empty_reply")


def test_list_items_render_bold_and_absolute_value() -> None:
    """Numbered misconceptions keep bold labels and drop math delimiters."""
    source = (
        "**一、核心迷思概念识别**\n\n"
        '1. **迷思三**："$|a| = a$是恒成立的"\n'
        "* **追问**：$|5|$ 是多少？$|-5|$ 是多少？\n"
        "### 场景一：打破绝对值迷思\n"
    )
    document = _open(source)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "**" not in text
    assert "$" not in text
    assert "###" not in text
    assert "一、核心迷思概念识别" in text
    assert "|a| = a" in text
    assert "|5|" in text
    assert "|-5|" in text
    assert "场景一：打破绝对值迷思" in text
    bold = [run.text for paragraph in document.paragraphs for run in paragraph.runs if run.bold]
    assert "迷思三" in bold
    assert "追问" in bold
    math = [run.text for paragraph in document.paragraphs for run in paragraph.runs if run.italic and "|" in run.text]
    assert "|a| = a" in math


def test_multiline_display_math_drops_delimiters() -> None:
    """A fenced formula is centered text, not raw dollar signs."""
    document = _open("$$\n|a| = a\n$$\n")
    paragraph = document.paragraphs[0]
    assert paragraph.text == "|a| = a"
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "$$" not in paragraph.text


def test_latex_commands_become_readable_text() -> None:
    """Fractions, powers, and roots survive as plain math text."""
    assert linearize_latex(r"\frac{1}{2}") == "1/2"
    assert linearize_latex(r"x^2") == "x²"
    assert linearize_latex(r"\sqrt{x}") == "√x"
    assert linearize_latex(r"a_{1}") == "a₁"
    text = _document_text(r"面积是 $\frac{1}{2}bh$。")
    assert "1/2" in text
    assert "frac" not in text
    assert "$" not in text


def test_ordered_lists_restart_and_links_keep_urls() -> None:
    """Each numbered list keeps the author's numbers, and link targets stay visible."""
    text = _document_text("1. 甲\n2. 乙\n\n1. 丙\n\n见 [说明](https://example.com/a)\n- [x] 已做\n")
    assert text.count("1.") == 2
    assert "2." in text
    assert "https://example.com/a" in text
    assert "](" not in text
    assert "[x]" in text
    assert "已做" in text


def test_bracket_display_math_spans_lines() -> None:
    """A multi-line \\[ \\] block is one formula, not leftover brackets."""
    text = _document_text("\\[\n|a|=a\n\\]\n后文还在\n")
    assert "|a|=a" in text
    assert "后文还在" in text
    assert "[" not in text
    assert "]" not in text


def test_empty_math_block_is_rejected() -> None:
    """A delimiter pair with no formula is not archived."""
    try:
        build_markdown_docx("$$\n$$\n")
    except ValueError as exc:
        assert str(exc) == "empty_reply"
    else:
        raise AssertionError("expected empty_reply")


def test_table_quote_and_code_are_structured() -> None:
    """Tables, quotes, and fenced code are not left as raw markdown."""
    source = "> 引用一句\n\n| 概念 | 例子 |\n| --- | --- |\n| 绝对值 | $|a|$ |\n\n```\ncode_sample\n```\n"
    text = _document_text(source)
    assert "引用一句" in text
    assert "概念" in text
    assert "绝对值" in text
    assert "|a|" in text
    assert "code_sample" in text
    assert "```" not in text
    assert "| --- |" not in text
